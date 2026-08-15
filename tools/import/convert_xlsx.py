#!/usr/bin/env python3
"""
Convert a season price-list workbook into products.csv (+ optional images/) for import_season.py.

    # 3S sheet, photos embedded in the workbook
    python3 convert_xlsx.py ~/Downloads/3s_2026-27_Price_List.xlsx --brand 3S

    # Trophyverse: no category column in the sheet, photos supplied separately
    python3 convert_xlsx.py ~/Downloads/TrophyVerse_2026-27_Price_List.xlsx \
        --brand Trophyverse --category Trophies --no-images

    # Verbose: tiered pricing, photos supplied separately
    python3 convert_xlsx.py ~/Downloads/Verbose_2026-27_Price_List.xlsx \
        --brand Verbose --no-images --append

Columns are matched by header name, so the three 2026-27 sheet layouts all work:
    item number   Item # | S. No.
    category      Category            (optional; use --category to set one for all rows)
    name          Product Name | Product
    specs         Specifications
    price         Price (Rs.) | Base Rate (Rs.)
    moq           MOQ                 (else parsed out of the specs text)
    packing       Master Packing      (else parsed out of the specs text)
"""

import argparse
import csv
import pathlib
import re
import sys
import unicodedata

try:
    import openpyxl
except ImportError:
    sys.exit("openpyxl is required:  pip3 install openpyxl")

HERE = pathlib.Path(__file__).parent

# Sheet category name -> category name used in the app's Category class.
CATEGORY_MAP = {
    "Table Clocks/Desk Organizer": "Table Clocks",
}

HEADER_ALIASES = {
    "item":     ("item #", "s. no.", "s.no.", "sr. no.", "item no"),
    "category": ("category",),
    "name":     ("product name", "product", "name"),
    "specs":    ("specifications", "specification", "details", "detail"),
    "price":    ("price (rs.)", "base rate (rs.)", "price", "rate"),
    "moq":      ("moq", "minimum qty"),
    "packing":  ("master packing", "master pack"),
}


def slug(text):
    text = unicodedata.normalize("NFKD", str(text)).encode("ascii", "ignore").decode()
    return re.sub(r"[^a-z0-9]+", "-", text.lower()).strip("-") or "item"


def map_headers(ws):
    header = [(str(c).strip().lower() if c is not None else "")
              for c in next(ws.iter_rows(min_row=1, max_row=1, values_only=True))]
    found = {}
    for field, aliases in HEADER_ALIASES.items():
        for idx, cell in enumerate(header):
            if cell in aliases:
                found[field] = idx
                break
    missing = [f for f in ("name", "price") if f not in found]
    if missing:
        sys.exit("Could not find column(s) %s in header: %s" % (missing, header))
    return found


def number_from_text(text, *labels):
    for label in labels:
        m = re.search(label + r"\s*:?\s*([\d,]+)", text or "", re.IGNORECASE)
        if m:
            return int(m.group(1).replace(",", ""))
    return None


def to_number(value):
    if value is None:
        return None
    if isinstance(value, (int, float)):
        return int(value) if float(value).is_integer() else value
    digits = re.sub(r"[^\d.]", "", str(value))
    if not digits:
        return None
    number = float(digits)
    return int(number) if number.is_integer() else number


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("workbook")
    ap.add_argument("--brand", required=True, help="3S | Trophyverse | Verbose")
    ap.add_argument("--sheet", default=None)
    ap.add_argument("--category", default=None,
                    help="force this category for every row (sheets with no category column)")
    ap.add_argument("--no-images", action="store_true",
                    help="skip embedded workbook images (photos supplied separately)")
    ap.add_argument("--append", action="store_true", help="add to products.csv instead of replacing it")
    args = ap.parse_args()

    wb = openpyxl.load_workbook(args.workbook, data_only=True)
    ws = wb[args.sheet] if args.sheet else wb.active
    col = map_headers(ws)

    images_by_row = {}
    if not args.no_images:
        for image in getattr(ws, "_images", []):
            images_by_row[image.anchor._from.row + 1] = image
        (HERE / "images").mkdir(exist_ok=True)

    rows_out, no_image, categories = [], [], set()

    for row_no, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        cell = lambda field: (row[col[field]] if field in col and col[field] < len(row) else None)
        name = cell("name")
        if not name or not str(name).strip():
            continue

        specs = str(cell("specs") or "").strip()
        sheet_category = args.category or (cell("category") or "")
        category = CATEGORY_MAP.get(str(sheet_category).strip(), str(sheet_category).strip())
        if category:
            categories.add(category)

        item_no = to_number(cell("item")) or row_no - 1

        image_name = ""
        if not args.no_images:
            image = images_by_row.get(row_no)
            if image is None:
                no_image.append((row_no, name))
            else:
                image_name = "%03d-%s.png" % (item_no, slug(name))
                (HERE / "images" / image_name).write_bytes(image._data())

        rows_out.append({
            "Name": str(name).strip(),
            "Brand": args.brand,
            "Category": category,
            "Price": to_number(cell("price")) or "",
            "Detail": specs,
            "Serial": "",
            "MasterSerial": "",
            "Order": item_no,
            "MOQ": to_number(cell("moq")) or number_from_text(specs, "Minimum Qty", "MOQ", "Min Qty") or "",
            "MasterPacking": to_number(cell("packing")) or number_from_text(specs, "Master Packing", "Master Pack") or "",
            "isNew": "True",
            "image_file": image_name,
        })

    out = HERE / "products.csv"
    existing = []
    if args.append and out.exists():
        with out.open(newline="", encoding="utf-8") as fh:
            existing = list(csv.DictReader(fh))

    with out.open("w", newline="", encoding="utf-8") as fh:
        writer = csv.DictWriter(fh, fieldnames=list(rows_out[0].keys()))
        writer.writeheader()
        writer.writerows(existing + rows_out)

    print("%s %d %s products -> products.csv (%d rows total)"
          % ("Appended" if existing else "Wrote", len(rows_out), args.brand, len(existing) + len(rows_out)))
    if categories:
        print("Categories used (%d): %s" % (len(categories), ", ".join(sorted(categories))))
    missing_price = [r["Name"] for r in rows_out if r["Price"] == ""]
    if missing_price:
        print("Rows with no price (%d): %s" % (len(missing_price), ", ".join(missing_price[:6])))
    if no_image:
        print("Rows with no embedded image: %d" % len(no_image))
    if args.no_images:
        print("Images skipped — attach them later with attach_images.py")


if __name__ == "__main__":
    main()
